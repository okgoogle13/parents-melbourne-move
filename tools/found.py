import openpyxl, re, sys
OLD, WB, OUT = sys.argv[1:4]
old = open(OLD, encoding='utf-8').read()
ws = openpyxl.load_workbook(WB, data_only=True)['SuburbScorecard']

W = {}
for r in range(2, 73):
    a = ws[f'A{r}'].value
    if not a or ws[f'C{r}'].value == 'Synthetic Test':
        continue
    d = {c: ws[f'{c}{r}'].value for c in
         ['A', 'C', 'G', 'H', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V', 'X', 'Y', 'AB', 'AD', 'AQ', 'AU', 'AX']}
    d['row'] = r
    W[a] = d
RANKED = sorted([d for d in W.values() if d['Y'] is not None], key=lambda d: d['Y'])
assert len(RANKED) == 51, len(RANKED)


def basis(d):
    if d['C'] == 'Community comparator':
        return '10:00 weekday, 28 Oct'
    return 'Leave now, 27 Oct' if '27 Oct' in str(d['AD']) else 'Leave now, 26 Oct'


def sc(d):
    return f"{d['X']:.2f} (#{d['Y']})"


OPT = {'Frankston': 'A', 'Seaford': 'AB', 'Mordialloc': 'C'}
WORD = {'Nature': 'N', 'Village': 'R', 'Beach': 'U', 'Shops': 'Q', 'Indian': 'P'}


def fix_scores(text, d):
    def rep(m):
        v = d[WORD[m.group(1)]]
        return f"{m.group(1)} {v:g}"
    return re.sub(r'\b(Nature|Village|Beach|Shops|Indian) (\d+)\b', rep, text)


def sect(start, end):
    i = old.index(start)
    j = old.index(end, i)
    return old[i:j]


# ---------- master table: reuse region, price string, train, hazard from the old table; numbers from the workbook
master_old = sect('| **DANDENONG FOOTHILLS AND RIDGE**', '\n\n')
mrows = []
seen = set()
for line in master_old.strip().split('\n'):
    cells = [c.strip() for c in line.strip('|').split('|')]
    if cells[0].startswith('**'):
        mrows.append(f"| {cells[0]} | {cells[1]} | | | | | | | | | | | | | |")
        continue
    name = cells[0]
    d = W[name]
    seen.add(name)
    gw_old, ed_old = cells[4], re.sub(r'[!?]', '', cells[5])
    assert int(gw_old) == d['J'], (name, gw_old, d['J'])
    assert int(ed_old) == d['K'], (name, ed_old, d['K'])
    tr = d['AB'][0] if d['AB'] else '?'
    assert tr == cells[6], (name, tr, cells[6])
    mrows.append('| ' + ' | '.join([
        name, cells[1], OPT.get(name, '–'), cells[3], str(d['J']), basis(d), cells[5], tr,
        f"{d['N']:g}", f"{d['R']:g}", f"{d['U']:g}", f"{d['Q']:g}", f"{d['P']:g}", sc(d), cells[13]]) + ' |')
assert seen == {d['A'] for d in RANKED}, seen ^ {d['A'] for d in RANKED}
MASTER = ('| Suburb | Rg | Opt | House price & basis | GW min | GW drive basis | ED min | Train | Nat | Vil | Bch | Shp | Ind | Workbook score (rank) | Hazard |\n'
          '|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|\n' + '\n'.join(mrows))

TOP = '| Rank | Suburb | Score | GW min (basis) | ED min | Nat | Vil | Bch | Shp | Ind | Train score |\n|---|---|---|---|---|---|---|---|---|---|---|\n'
for d in RANKED[:15]:
    TOP += f"| {d['Y']} | {d['A']} | {d['X']:.2f} | {d['J']} ({basis(d)}) | {d['K']} | {d['N']:g} | {d['R']:g} | {d['U']:g} | {d['Q']:g} | {d['P']:g} | {d['S']:g} |\n"

# ---------- cards
cards_old = sect('**Upwey** · DAN', '### The other 30 passing suburbs')
cards = {}
for blk in re.split(r'\n(?=\*\*[^*]+\*\* · )', cards_old.strip()):
    blk = re.sub(r'\n+### .*$', '', blk.strip(), flags=re.S).strip()
    head, body = blk.split('\n', 1)
    name = re.match(r'\*\*([^*]+)\*\*', head).group(1)
    base = re.sub(r' \((control)\)', '', name)
    d = W[base]
    parts = head.split(' · ')
    # parts: name, region, options, price, GW, ED, train, score
    gw = int(parts[4].split()[1]); ed = int(re.sub(r'\D', '', parts[5].split()[1]))
    assert gw == d['J'] and ed == d['K'], (base, gw, ed)
    newhead = ' · '.join([f"**{base}**", parts[1], parts[3], f"GW {d['J']} ({basis(d)})", parts[5], parts[6], f"workbook {sc(d)}"])
    body = fix_scores(body, d)
    body = body.replace('**Estimate:** ', '**Estimate (workbook scores and notes):** ')
    cards[base] = newhead + '\n' + body.strip()

cards['Croydon'] = cards['Croydon'].replace(
    '**Role:** shows the cheapest "village plus train" option with modest nature; if the family prefers it, nature is not their real priority.',
    '**Trade-off:** a cheap "village plus train" option 18 minutes from family, with modest nature.')
cards['Berwick'] = cards['Berwick'].replace(
    '**Role:** the control for Q13, showing what a strong Indian community and services feel like with little forest and no beach.',
    '**Trade-off:** the highest India-born share of the ranked top 10 (6.1%) with strong services; little forest and no beach.')

ORDER = {'Dandenong foothills and ridge': ['Upwey', 'Belgrave', 'Upper Ferntree Gully', 'Sassafras', 'The Basin', 'Mount Evelyn'],
         'North-east and inner-east': ['Eltham', 'Diamond Creek', 'Warrandyte', 'Hurstbridge', 'Croydon'],
         'Bayside, south-east and Peninsula': ['Berwick', 'Frankston South', 'Mornington', 'Mount Martha', 'Dromana'],
         'Macedon, Geelong and Gippsland': ['Woodend', 'Highton (Geelong)', 'Torquay', 'Warragul']}
used = sum(ORDER.values(), []) + ['Seaford']
assert set(used) == set(cards), set(cards) ^ set(used)
OTHER_CARDS = ''
for g, names in ORDER.items():
    OTHER_CARDS += f"### {g}\n\n" + '\n\n'.join(cards[n] for n in names) + '\n\n'

# ---------- one-liners (Frankston and Mordialloc now have cards)
ol = sect('| Suburb (rank) | Note |', '\n\n## Where')
orows = []
for line in ol.strip().split('\n')[2:]:
    cells = [c.strip() for c in line.strip('|').split('|')]
    name = re.sub(r' \(#\d+\)', '', cells[0])
    key = {'Newtown, Geelong': 'Newtown (Geelong)'}.get(name, name)
    if key in ('Frankston', 'Mordialloc'):
        continue
    d = W[key]
    note = fix_scores(cells[1], d)
    orows.append((d['Y'], f"| {name} (#{d['Y']}, {d['X']:.2f}) | {note} |"))
orows.sort()
ONELINE = '| Suburb (rank, score) | Note |\n|---|---|\n' + '\n'.join(r for _, r in orows)
assert len(orows) == 51 - 3 - len(used) + 1, len(orows)  # 51 ranked minus 3 anchor cards minus 20 other cards

# ---------- sections reused with edits
disagree = sect('| # | Item | Workbook says |', '\n\n## What is still')
unver = sect('| Gap | What is known |', '\n\n### At inspection')
inspect = sect('| Check | Why |', '\n\n### What Round 2')
sources = sect('| Topic | Source |', '\n| Community comparators |')
srcrow = old[old.index('| Community comparators |'):].strip()

tpl = open(sys.argv[4], encoding='utf-8').read()
out = (tpl.replace('{{MASTER}}', MASTER).replace('{{TOP}}', TOP.strip()).replace('{{SEAFORD}}', cards['Seaford'])
       .replace('{{OTHER_CARDS}}', OTHER_CARDS.strip()).replace('{{ONELINE}}', ONELINE)
       .replace('{{DISAGREE}}', disagree.strip()).replace('{{UNVER}}', unver.strip()).replace('{{INSPECT}}', inspect.strip())
       .replace('{{SOURCES}}', sources.strip() + '\n' + srcrow))
for k in ['Frankston', 'Mordialloc', 'Seaford']:
    out = out.replace('{{%s}}' % k, sc(W[k]))
assert '{{' not in out, re.findall(r'\{\{\w+\}\}', out)
open(OUT, 'w', encoding='utf-8').write(out)
print('ok', len(out))
