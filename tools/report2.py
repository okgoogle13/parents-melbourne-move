import openpyxl, html, re, sys
WB=sys.argv[1]; OUT=sys.argv[2]
wb=openpyxl.load_workbook(WB,data_only=True); ws=wb['SuburbScorecard']
E=html.escape
COLS={'A':1,'B':2,'C':3,'D':4,'E':5,'G':7,'H':8,'J':10,'K':11,'L':12,'M':13,'N':14,'O':15,'P':16,'Q':17,'R':18,'S':19,'T':20,'U':21,'V':22,'X':24,'Y':25,'AD':30,'AG':33,'AI':35,'AJ':36,'AQ':43,'AU':47,'AX':50}
rows=[]
for r in range(2,73):
    if not ws.cell(r,1).value or ws.cell(r,3).value=='Synthetic Test': continue
    d={k:ws.cell(r,c).value for k,c in COLS.items()}; d['row']=r; rows.append(d)
by={d['A']:d for d in rows}
GROUP={'South / Coastal':1,'Bayside Melbourne':1,'Mornington Peninsula':2,'Dandenong Ranges':3,'Dandenong Ridge':3,'Dandenong Foothills':3,'East / Foothills':3,
 'East Melbourne':4,'Outer East Melbourne':4,'North-East Melbourne':4,'North-East':4,'North-East / Yarra':4,'North-East Fringe':4,'South-East Melbourne':4,
 'Yarra Valley':5,'Yarra Ranges':5,'West':6,'West Coastal':6,'Geelong':6,'Bellarine Peninsula':6,'Surf Coast':6,
 'Macedon Ranges':7,'Central Victoria':7,'Central Highlands':7,'West Gippsland':7,'Bass Coast':7}
GN={1:'Bayside / South coast',2:'Mornington Peninsula',3:'Dandenong Ranges & Foothills',4:'East, North-East & South-East Melbourne',5:'Yarra Valley & Ranges',6:'West / Geelong / Bellarine / Surf Coast',7:'Regional towns'}
for d in rows: d['grp']=GN[GROUP[d['B']]]
COMP=['Springvale','Clayton South','Keysborough','Clyde North','Endeavour Hills','Wheelers Hill']
def is_comp(d): return d['C']=='Community comparator'
# ---- evidence labels (derived from workbook text only)
def tag(k):
    return {'C':'<span class="ev ev-c">Checked</span>','E':'<span class="ev ev-e">Est.</span>','V':'<span class="ev ev-v">Verify</span>','NA':'<span class="ev ev-na">N/A</span>'}[k]
def price_ev(d):
    h=str(d['H']) 
    if is_comp(d): return 'C'
    if 'AVM' in h or 'NOT a verified' in h or 'only 5 house sales' in h: return 'E'
    return 'C'
def price_basis(d):
    h=str(d['H'])
    if is_comp(d): return 'Suburb sale median (Statement of Information)'
    if 'AVM' in h: return 'Automated valuation (not a sale median)'
    if 'only 5' in h: return 'Sale median, only 5 sales'
    m=re.match(r'(3-bed house median)',h)
    if m: return '3-bed house sale median (Domain)'
    return 'Suburb sale median'
def drive_basis(d):
    if is_comp(d): return '10:00 weekday, 28 Oct'
    ad=str(d['AD'])
    if '27 Oct' in ad: return 'Leave now, 27 Oct'
    return 'Leave now, 26 Oct'
def lc(x): return x[0].lower()+x[1:]
def drive_ev(d):
    return 'E' if 'Origin approximated' in str(d['AD']) else 'C'
def ed_name(d):
    ad=str(d['AD'])
    if is_comp(d): return ad.split(',')[0]
    return re.split(r' \(|\. ',ad)[0]
def ed_status(d):
    ad=str(d['AD'])
    if is_comp(d) or '24-hr: verified' in ad: return 'C'
    if 'still to verify' in ad or 'unverified' in ad or 'not confirmed' in ad or 'less serious' in ad: return 'V'
    return 'E'
def ed_note(d):
    ad=str(d['AD'])
    if is_comp(d): return 'public 24-hour ED, checked'
    if '24-hr: verified' in ad: return '24-hour status checked'
    if 'still to verify' in ad or 'unverified' in ad: return '24-hour status to verify'
    if 'not confirmed' in ad: return 'nearest ED not confirmed'
    if 'less serious' in ad: return 'ED treats less serious cases'
    if '24-hr: recall' in ad: return '24-hour status from memory'
    return '24-hour status not recorded'
def fil(d):
    l=d['L']
    if l in ('Active — provisional pass','Provisional pass'): return '<span class="st st-pp">Provisional pass</span>'
    if l=='Active — needs data': return '<span class="st st-nd">Needs data</span>'
    if str(l).startswith('Fails'): return '<span class="st st-f">Fail</span> Glen Waverley drive'
    if is_comp(d): return 'Filters met (comparator, not ranked)'
    return E(str(l))
def money(v): return '$'+format(int(round(v)),',') if isinstance(v,(int,float)) else '—'
def pct(v): return f'{v*100:.1f}%' if isinstance(v,(int,float)) else '—'
def num(v,dp=1):
    if v is None or v=='': return '<span class="blank">—</span>'
    return f'{v:.{dp}f}' if isinstance(v,float) else str(v)
def wsc(d): return f"{d['X']:.2f}" if isinstance(d['X'],(int,float)) else '<span class=na>N/A</span>'
def rank(d): return str(d['Y']) if isinstance(d['Y'],(int,float)) else 'N/A'
def issues(d):
    out=[]
    if is_comp(d):
        out.append('Suitable-property check outstanding')
        if d['A']=='Endeavour Hills': out.append('Endeavour Spices and Shri Sanatan Dharam need verification')
        return '; '.join(out)
    if d['L']=='Active — needs data': return E(str(d['AJ']).split('. ')[0])+'.'
    if str(d['L']).startswith('Fails'): return f"Glen Waverley drive {d['J']} min is over the 90-min limit"
    if price_ev(d)=='E': out.append('Price is an automated valuation' if 'AVM' in str(d['H']) else 'Price based on few sales')
    n=ed_note(d)
    if ed_status(d)!='C': out.append(n[0].upper()+n[1:])
    if drive_ev(d)=='E': out.append('Drive origin approximated')
    return '; '.join(out) if out else 'Suitable home not yet checked'
# ---- comparator facilities (research file, approval pack A)
FAC={
 'Springvale':('10–26','India Town, 4 min','C','Sankat Mochan Hanuman Mandir (Oakleigh South), 10–20 min','C',"St Joseph's, 4 min",'None of real size named'),
 'Clayton South':('14–35','Sarawan Spices, 5–8 min','C','Sankat Mochan Hanuman Mandir, 7–12 min','C',"St Andrews, 5 min",'No named reserve'),
 'Keysborough':('18–24','India at Home (Parkmore), 4–6 min','C','Shri Shiva Vishnu (Carrum Downs), 14–20 min','C','Resurrection Parish, 5–6 min','Market gardens; no named reserve'),
 'Clyde North':('28–50','Wirsa Indian Groceries, 2 min','C','Shrinathji Haveli (Hallam), 16–30 min','C','St Thomas The Apostle, 4 min','Farmland; park only proposed'),
 'Endeavour Hills':('16–26','India At Home Dandenong, 8–12 min','C','Shrinathji Haveli (Hallam), 8–14 min','C','St Paul Apostle, 2 min','Bush reserve 1.1 km (3 min); Churchill National Park 8–10 min'),
 'Wheelers Hill':('7–12','The Heritage Spices, 6–9 min','C','Shree Swaminarayan (Boronia), 10–20 min','C',"St Justin's, 4 min",'Jells Park 2.6 km (6 min)'),
}
PEAK={'Springvale':'10–35','Clayton South':'16–40','Keysborough':'18–30','Clyde North':'35–70','Endeavour Hills':'18–35','Wheelers Hill':'4–12'}
for _n,_f in FAC.items(): assert int(_f[0].split('–')[1])==by[_n]['J'],(_n,_f[0],by[_n]['J'])
CORNER={'Springvale':'Comparator','Clayton South':'Comparator','Keysborough':'Comparator','Clyde North':'Comparator','Endeavour Hills':'Comparator — Potential both','Wheelers Hill':'Comparator — Potential both (community lighter)'}
# ---- weights
fc=wb['ForcedChoiceCorrections']
W=sorted([(fc.cell(r,2).value,fc.cell(r,1).value) for r in range(6,16)],reverse=True)
PLAIN={'Near family in GW':'Near family in Glen Waverley','Village character':'Village feel','Train/bus':'Train or bus','Shops & doctor':'Shops & doctor','City trip':'Trip to the city'}
ds=wb['DiscussionSummary']
filt=wb['Filters']
cap=filt['B2'].value; gwmax=filt['B4'].value; edmax=filt['B5'].value; outlay=filt['B3'].value
active=[d for d in rows if d['C']=='Active Shortlist']
assert len(active)==32,len(active)
A_pp=sorted([d for d in active if d['L']=='Active — provisional pass'],key=lambda d:d['Y'])
A_nd=[d for d in active if d['L']=='Active — needs data']
A_f=[d for d in active if str(d['L']).startswith('Fails')]
assert len(A_pp)+len(A_nd)+len(A_f)==32
TOPN=int(sys.argv[3]) if len(sys.argv)>3 else 8
CSS='''
:root{--ink:#1a1a1a;--mute:#555;--rule:#bbb;--bg:#fff;--pp:#e3f1e3;--ppb:#2e6b2e;--cp:#1f5d1f;--nd:#fbe9c6;--ndb:#8a5a00;--f:#f6d4d0;--fb:#9b1c12;--soft:#f3f3f3}
@page{size:A4 portrait;margin:10mm 11mm}
@page land{size:A4 landscape;margin:10mm}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:9.4pt/1.3 "Helvetica Neue",Arial,sans-serif}
h1{font-size:15pt;margin:0 0 2mm} h2{font-size:11pt;margin:2.2mm 0 1mm;border-bottom:1px solid var(--rule);padding-bottom:.6mm} h3{font-size:9.8pt;margin:2mm 0 1mm}
p{margin:0 0 1.4mm} ul{margin:0 0 1.4mm;padding-left:4.5mm} li{margin:0 0 .5mm}
.page{page-break-after:always;break-after:page}
.land{page:land}
.box{border:1.2px solid var(--ink);padding:2mm 3mm;margin:0 0 2mm;break-inside:avoid}
.box p{margin:0 0 1.2mm}
.key{display:flex;flex-wrap:wrap;gap:1.2mm 3mm;font-size:8.2pt;margin:0 0 2mm;break-inside:avoid}
.st{display:inline-block;border:1px solid;border-radius:2px;padding:0 1.2mm;font-size:7.8pt;font-weight:600;white-space:nowrap}
.st-pp{background:var(--pp);border-color:var(--ppb);color:var(--ppb)} .st-cp{background:var(--cp);border-color:var(--cp);color:#fff}
.st-nd{background:var(--nd);border-color:var(--ndb);color:var(--ndb)} .st-f{background:var(--f);border-color:var(--fb);color:var(--fb)}
.ev{font-size:7.2pt;white-space:nowrap;margin-left:.8mm} .ev-c{color:var(--ink)} .ev-e{color:var(--mute);font-style:italic;border:1px dotted var(--mute);padding:0 .6mm}
.ev-v{color:var(--ndb);border:1px solid var(--ndb);padding:0 .6mm;font-weight:600} .ev-na,.na{color:var(--mute)}
table{border-collapse:collapse;width:100%;margin:0 0 2mm;font-size:8.1pt} thead{display:table-header-group}
th,td{border:1px solid var(--rule);padding:.5mm 1mm;vertical-align:top;text-align:left} th{background:var(--soft);font-weight:600}
tr{break-inside:avoid} td.nw,span.nw{white-space:nowrap}td.n{text-align:right;white-space:nowrap}
.small{font-size:7.8pt;color:var(--mute)} .blank{color:var(--mute)}
.block{break-inside:avoid}
.cols{display:grid;grid-template-columns:1fr 1fr;gap:3mm}
.land table{font-size:7.3pt} .land th,.land td{padding:.4mm .8mm}
table.nvc{table-layout:fixed}
.oq .block{column-count:2;column-gap:5mm}.oq ul{break-inside:avoid}
.draft{display:inline-block;border:1.5px solid #9b1c12;color:#9b1c12;font-weight:700;padding:.2mm 1.5mm;font-size:7.6pt;margin:0 0 .8mm}
section.sheet:not(.land) h2{margin:1.6mm 0 .8mm}
table.nvc td{padding:.3mm .8mm}
.draftsm{color:#9b1c12;font-weight:700;font-size:7.8pt;margin:0 0 1mm}
.note{font-size:8.1pt;border-left:3px solid var(--ink);padding:1mm 2.5mm;margin:0 0 2mm;break-inside:avoid}
@media screen{body{background:#ddd}.sheet{background:#fff;width:210mm;min-height:297mm;margin:6mm auto;padding:12mm;box-shadow:0 1px 4px #0003}.sheet.land{width:297mm;min-height:210mm;padding:10mm}}
'''
def ev_cell(val,k,basis=''):
    return f'{val} {tag(k)}'+(f'<br><span class="small">{E(basis)}</span>' if basis else '')
H=[]
H.append(f'<!doctype html><html lang="en-AU"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Draft: Our Next Home report</title><style>{CSS}</style></head><body>')
# ===== PAGE 1
H.append('<section class="sheet page"><div class="draft">Draft — requires family review; not for publication or distribution.</div><h1>Our Next Home: where we have got to (28 Oct 2026)</h1>')
H.append('''<div class="box"><p><b>What this is.</b> A summary of the suburbs we have looked at, scored against what the three of us said matters most. It is a research summary, not a decision.</p>
<p><b>How to read it.</b> Every place must meet our hard limits on price, drive time to Glen Waverley and drive time to a hospital emergency department. &ldquo;Provisional pass&rdquo; means a place meets them on current evidence; no place is a &ldquo;Confirmed pass&rdquo; yet, because we have not checked any specific home. Each fact is labelled <i>Checked</i> (we confirmed it from a named source), <i>Estimate</i> (best available figure, not confirmed) or <i>Needs verification</i> (not yet confirmed). Scores run from 0 to 10, where 10 is the best fit for that priority.</p>
<p><b>No place in this report is described as suitable</b> until both its hard limits and a specific home have been checked.</p></div>''')
H.append('<div class="key"><b>Labels:</b> <span class="ev ev-c">Checked</span> <span class="ev ev-e">Est.</span>&nbsp;Estimate <span class="ev ev-v">Verify</span>&nbsp;Needs verification <span class="st st-pp">Provisional pass</span> <span class="st st-cp">Confirmed pass</span>&nbsp;(no place yet) <span class="st st-f">Fail</span> <span class="st st-nd">Needs data</span> <span class="ev ev-na">N/A</span>&nbsp;Not applicable</div>')
H.append(f'''<h2>Our hard limits</h2><table><thead><tr><th>Limit</th><th>Value</th><th>Note</th></tr></thead><tbody>
<tr><td>Purchase price</td><td>{money(cap)} maximum</td><td>Total with 6% costs: {money(outlay)}</td></tr>
<tr><td>Drive to Glen Waverley</td><td>{gwmax} min maximum</td><td>60 min preferred: a preference, not a limit</td></tr>
<tr><td>Drive to a public 24-hour emergency department (ED)</td><td>{edmax} min maximum</td><td>ED drive time is also scored</td></tr>
<tr><td>Stairs</td><td><span class="na">Not applicable</span></td><td>Stairs are allowed; not a hard limit</td></tr>
<tr><td>Train</td><td><span class="na">Not applicable</span></td><td>Preferred, not required; scored instead</td></tr></tbody></table>''')
H.append('<div class="cols"><div class="block"><h2>What matters most to us</h2><ol style="margin:0 0 1mm;padding-left:5mm">')
for w,n in W: H.append(f'<li>{E(PLAIN.get(n,n))} — {round(w*100)}%</li>')
H.append('</ol><p class="small">Each place gets a 0–10 score for each priority; the weighted score combines them, out of 10.</p></div>')
H.append(f'''<div class="block"><h2>Where we differ</h2><p><b>Distance and setting</b> (survey questions 13b and 13c):</p><ul>
<li>Mum: {E(str(ds['E23'].value))}; {E(str(ds['F23'].value))}</li>
<li>Dad: {E(str(ds['E24'].value))}; {E(str(ds['F24'].value))}</li>
<li>Me: {E(str(ds['E25'].value))}; {E(str(ds['F25'].value))}</li></ul>
<p><b>Still to discuss:</b> living arrangement ({E(str(ds['E17'].value).split(' (notes')[0])}) and garden ({E(str(ds['E18'].value))}).</p></div></div></section>''')
H[-1]=H[-1].replace('</section>','')
top=A_pp[:TOPN]; tb=sorted(set(drive_basis(d) for d in top))
H.append(f'<h2>Top {TOPN} active locations — Provisional pass</h2><table><thead><tr><th>Rank</th><th>Location</th><th>Report group</th><th>Weighted score (0–10)</th><th>Price (basis)</th><th>Glen Waverley drive ({"; ".join(lc(b) for b in tb)})</th><th>ED drive (24-hour status)</th></tr></thead><tbody>')
for d in top:
    pb='sale median' if price_ev(d)=='C' else 'automated valuation'
    H.append(f"<tr><td class=n>{rank(d)}</td><td>{E(d['A'])}</td><td>{E(d['grp'])}</td><td class=n>{d['X']:.2f}</td><td class=nw>{money(d['G'])} {tag(price_ev(d))} <span class=small>{pb}</span></td><td>{d['J']} min {tag(drive_ev(d))}</td><td>{d['K']} min {tag(ed_status(d))}</td></tr>")
H.append('</tbody></table><p class="small">Rank is among all 51 provisional-pass places, including Watchlist places; gaps are Watchlist places. ED label = whether the hospital\'s 24-hour public ED status is checked (drive minutes are from Google Maps). All 32 active locations: Appendix A3.</p>')
H.append('<div class="cols"><div class="block"><h3><span class="st st-nd">Needs data</span> active locations</h3><ul>')
ND={'Blackburn':'house median $1.60M (sale median) is over the cap; unit median $820k',
    'Glen Waverley':'house median $1.76M (sale median) is over the cap; unit median $930k',
    'Williamstown':'$1.64M is an automated valuation, not a verified sale median',
    'Mount Eliza':'$1.61M automated valuation is over the cap; not a verified sale median'}
for d in A_nd: H.append(f"<li><b>{E(d['A'])}</b>: {ND[d['A']]}.</li>")
H.append('</ul></div><div class="block"><h3><span class="st st-f">Fail</span> active locations</h3><ul>')
for d in A_f: H.append(f"<li><b>{E(d['A'])}</b>: {d['J']} min to Glen Waverley ({lc(drive_basis(d))}, not a peak-hour result) — over the 90-min limit. Price {money(d['G'])} passes the cap.</li>")
H.append('</ul></div></div>')
# ===== PAGE 2 (flows on from page 1)

# nature vs community
def nvc_row(name,role,gw,price,ed,aq,au,band,fac):
    return f'<tr><td><b>{name}</b><br><span class=small>{role}</span></td><td>{gw}</td><td>{price}</td><td>{ed}</td><td class=n>{aq} / {au}</td><td>{band}</td><td>{fac}</td></tr>'
H.append('<h2>Nature versus community</h2><table class="nvc"><colgroup><col style="width:17%"><col style="width:12%"><col style="width:16%"><col style="width:13%"><col style="width:11%"><col style="width:7%"><col style="width:24%"></colgroup><thead><tr><th>Location and role</th><th>Glen Waverley drive</th><th>Price</th><th>ED drive</th><th>India-born / other language at home</th><th>Band</th><th>Nearest Indian grocery / Hindu temple (min; all Checked)</th></tr></thead><tbody>')
for n,role in (('Frankston','Anchor (Active)'),('Mordialloc','Anchor (Watchlist)'),('Seaford','Anchor (Watchlist)')):
    d=by[n]; pb='sale median' if price_ev(d)=='C' else 'auto. valuation'
    H.append(nvc_row(n,role,f"{d['J']} min {tag(drive_ev(d))}",f"{money(d['G'])} {tag(price_ev(d))}",f"{d['K']} min {tag(ed_status(d))}",pct(d['AQ']),pct(d['AU']),E(str(d['AX'])),'<span class=na>N/A</span>'))
SHORT={'Springvale':('India Town 4','Sankat Mochan 10–20'),'Clayton South':('Sarawan Spices 5–8','Sankat Mochan 7–12'),'Keysborough':('India at Home 4–6','Shri Shiva Vishnu 14–20'),
 'Clyde North':('Wirsa 2','Shrinathji Haveli 16–30'),'Endeavour Hills':('India At Home (Dandenong) 8–12','Shrinathji Haveli 8–14'),'Wheelers Hill':('Heritage Spices 6–9','Swaminarayan (Boronia) 10–20')}
ROLE={'Endeavour Hills':'Comparator, Potential both','Wheelers Hill':'Comparator, Potential both (community lighter)'}
for n in COMP:
    d=by[n]; f=FAC[n]; g,t_=SHORT[n]
    H.append(nvc_row(n,ROLE.get(n,'Comparator'),f"{d['J']} min {tag('C')}",f"{money(d['G'])} {tag('C')}",f"{d['K']} min {tag('C')}",pct(d['AQ']),pct(d['AU']),E(str(d['AX'])),f"{E(g)}<br>{E(t_)}"))
H.append(nvc_row('Rowville','Reference only',f"~15 min {tag('E')}",f"~$1,129,555 {tag('E')}",f"~15–20 min {tag('E')}",'3.2%','31.6%','Medium','<span class=na>N/A</span>'))
H.append('</tbody></table><p class="small">Drive basis (as in the workbook, Glen Waverley and ED): anchors &ldquo;leave now&rdquo; 26–27 Oct; comparators 10:00 weekday, 28 Oct (top of range). Comparator peak (08:00) ranges: A4, reference only. Prices: Frankston and comparators sale medians; Mordialloc, Seaford, Rowville automated valuations. Rowville figures not looked up; its ED is not confirmed.</p>')
H.append('<div class="note">This table puts two kinds of place side by side. <b>Nature-first anchors</b> (Frankston, Mordialloc, Seaford) are green and coastal, with smaller Indian communities. Mordialloc and Seaford are Watchlist locations, not on the active shortlist. <b>Community comparators</b> (Springvale, Clayton South, Keysborough, Clyde North, Endeavour Hills, Wheelers Hill) have larger Indian and multilingual communities but less nature. Endeavour Hills and Wheelers Hill have some of both. Comparators are <b>not ranked and carry no weight</b> in the main score; they are here only to make the trade-off between nature and community visible. The <b>Community comparison band</b> (High / Medium / Low) is based on census language and birthplace figures. <b>It is not a liveability score.</b> Rowville is shown for reference only. <b>Point Cook is Conditional, not a comparator:</b> it fails the verified 90-minute Glen Waverley journey at 08:00 peak, and including it would need a family decision to accept a longer travel trade-off.</div>')
OQ='''<div class="block"><p><b>Still to decide</b></p><ul>
<li>What the &ldquo;Indian community&rdquo; score should measure. This is pending a family decision.</li>
<li>How we will use the answers to the community-versus-nature question (Q13d). Reference-only; feeds no score or weight.</li>
<li>Whether a longer trip to Glen Waverley is acceptable for Point Cook. A family decision.</li>
<li>Whether our 60-minute preference for the Glen Waverley drive should become a hard limit. Not yet decided; today the limit is 90 minutes.</li></ul>
<p><b>Still to check</b></p><ul>
<li>Whether a suitable home is actually available in each comparator suburb. Not yet checked for any of them.</li>
<li>Train access for the comparator suburbs. Currently an estimate only.</li></ul>
<p><b>Next step:</b> agree these questions, then choose places to view. The proposed rule is at least 2 community comparators, at least 1 nature-first anchor, and at least 1 checked &ldquo;Potential both&rdquo; suburb.</p></div>'''
H.append('<h2>Open questions and next steps</h2><div class="oq">'+OQ+'</div></section>')
# ===== APPENDIX
def sec(title,body,last=False): H.append(f'<section class="sheet land{"" if last else " page"}"><p class="draftsm">Draft — requires family review; not for publication or distribution.</p><h2>{title}</h2>{body}</section>')
order=A_pp+A_nd+A_f+[by['Mordialloc'],by['Seaford']]+[d for d in rows if d['C']=='Open-mind viewing']+[by[n] for n in COMP]
order+= [d for d in rows if d not in order]
assert len(order)==len(rows)==69,len(rows)
t=['<p class="small">Appendix tables exclude the two synthetic test rows. 69 locations: 32 Active Shortlist, 29 Watchlist, 2 Open-mind viewing, 6 Community comparators. Rows: Active (provisional pass, needs data, fail), anchors, Open-mind viewing, comparators, then the rest of the Watchlist.</p>',
 '<table><thead><tr><th>Location</th><th>Report group</th><th>Workbook status</th><th>Hard-filter result</th><th>Price</th><th>Price basis</th><th>Glen Waverley drive (min)</th><th>Drive basis</th><th>ED drive (min)</th><th>ED (from workbook)</th><th>Rank</th><th>Main open issue</th></tr></thead><tbody>']
for d in order:
    t.append(f"<tr><td>{E(d['A'])}</td><td>{E(d['grp'])}</td><td>{E(d['C'])}</td><td>{fil(d)}</td><td class=n>{money(d['G'])} {tag(price_ev(d))}</td><td>{E(price_basis(d))}</td><td class=n>{d['J']} {tag(drive_ev(d))}</td><td>{drive_basis(d)}</td><td class=n>{d['K']} {tag(ed_status(d))}</td><td>{E(ed_name(d))}; {ed_note(d)}</td><td class=n>{rank(d)}</td><td>{issues(d)}</td></tr>")
t.append('</tbody></table>')
sec('A1. Full comparison table — filters and evidence',''.join(t))
SC=[('M','Near family'),('N','Nature'),('O','Hospital'),('P','Indian community'),('Q','Shops & doctor'),('R','Village'),('S','Train/bus'),('T','Price'),('U','Beach'),('V','City trip')]
t=['<p class="small">Scores run 0–10; 10 is the best fit for that priority. Near family, Hospital and Price are calculated from drive times and price. The other seven are judgements, each with a source in the workbook. Blank = not scored. Comparators have no weighted score or rank (Not applicable). The comparison band is display-only and not a liveability score.</p><table><thead><tr><th>Location</th>'+''.join(f'<th>{n}</th>' for _,n in SC)+'<th>Weighted score (0–10)</th><th>Rank</th><th>India-born</th><th>Other language at home</th><th>Comparison band</th></tr></thead><tbody>']
for d in order:
    ws_=f"{d['X']:.2f}" if isinstance(d['X'],(int,float)) else '<span class=na>N/A</span>'
    t.append(f"<tr><td>{E(d['A'])}</td>"+''.join(f"<td class=n>{num(d[k])}</td>" for k,_ in SC)+f"<td class=n>{ws_}</td><td class=n>{rank(d)}</td><td class=n>{pct(d['AQ'])}</td><td class=n>{pct(d['AU'])}</td><td>{E(str(d['AX'])) if d['AX'] else '<span class=na>N/A</span>'}</td></tr>")
t.append('</tbody></table>')
sec('A2. Full comparison table — scores',''.join(t))
def a3(lst,title):
    s=[f'<h3>{title} ({len(lst)})</h3><table><thead><tr><th>Location</th><th>Report group</th><th>Hard-filter result</th><th>Weighted score</th><th>Rank</th><th>Why it is here / what is missing</th></tr></thead><tbody>']
    for d in lst:
        why=issues(d)
        if d in A_f: why=f"{d['J']} min to Glen Waverley ({lc(drive_basis(d))}, not peak) — over the 90-min limit; price {money(d['G'])} passes the cap"
        elif d in A_nd: why=ND[d['A']]
        s.append(f"<tr><td>{E(d['A'])}</td><td>{E(d['grp'])}</td><td>{fil(d)}</td><td class=n>{wsc(d)}</td><td class=n>{rank(d)}</td><td>{why}</td></tr>")
    return ''.join(s)+'</tbody></table>'
sec('A3. Active Shortlist — all 32 locations',a3(A_pp,'A. Active — provisional pass')+a3(A_nd,'B. Active — needs data')+a3(A_f,'C. Active — hard-filter fail')+'<p class="small">Frankston is also a nature-first anchor (A5). No workbook status has been changed.</p>')
t=['<p class="small">Comparison locations, not ranked candidates. They carry no weight in the main score. The Community comparison band is not a liveability score. Their purpose is to make the nature-versus-community trade-off visible. Glen Waverley and ED drives: Google Maps typical range departing 10:00 on a weekday, checked 28 Oct 2026; the top of the range is the workbook figure (J, K). Peak (08:00), not used in the workbook J/K figures: shown in its own column for reference only.</p>',
 '<table><thead><tr><th>Location</th><th>Role</th><th>Glen Waverley drive, 10:00 weekday (workbook J)</th><th>Peak (08:00), not used in the workbook J/K figures</th><th>Price</th><th>ED, 10:00 weekday (workbook K)</th><th>Indian grocery</th><th>Hindu temple</th><th>Church</th><th>Nature</th><th>Band</th><th>Open items</th></tr></thead><tbody>']
OPEN={'Springvale':'Suitable home not checked; train access estimate','Clayton South':'Suitable home not checked; train access estimate','Keysborough':'Suitable home not checked; train access estimate',
 'Clyde North':'Peak (08:00) top of 70 min would exceed a 60-min preference if adopted (workbook figure 50 min); suitable home not checked','Endeavour Hills':'Endeavour Spices <span class="ev ev-v">Verify</span>; Shri Sanatan Dharam hours <span class="ev ev-v">Verify</span> (4–6 min <span class="ev ev-e">Est.</span>); suitable home not checked',
 'Wheelers Hill':'Budget: sold houses under $1.5M (Jellis Craig, Sep 2026) <span class="ev ev-c">Checked</span> — supports further investigation; does not confirm a suitable home is available'}
for n in COMP:
    d=by[n]; f=FAC[n]
    t.append(f"<tr><td>{n}</td><td>{CORNER[n]}</td><td>{d['J']} min (range {f[0]}) {tag('C')}</td><td>{PEAK[n]} min {tag('C')}</td><td>{money(d['G'])} {tag('C')}<br><span class=small>suburb sale median</span></td><td>{E(ed_name(d))}, {str(d['AD']).split(',')[1].split(';')[0].strip()} {tag('C')}</td><td>{E(f[1])} {tag(f[2])}</td><td>{E(f[3])} {tag(f[4])}</td><td>{E(f[5])} {tag('C')}</td><td>{E(f[6])}</td><td>{d['AX']}</td><td>{OPEN[n]}</td></tr>")
t.append('</tbody></table>')
sec('A4. Community comparators',''.join(t))
t=['<table><thead><tr><th>Location</th><th>Workbook status</th><th>Rank</th><th>Glen Waverley drive</th><th>Price</th><th>ED</th><th>India-born</th><th>Other language</th><th>Band</th><th>Nature score</th></tr></thead><tbody>']
for n in ('Frankston','Mordialloc','Seaford'):
    d=by[n]; t.append(f"<tr><td>{n}</td><td>{d['C']}</td><td class=n>{rank(d)}</td><td>{ev_cell(str(d['J'])+' min',drive_ev(d),drive_basis(d))}</td><td>{ev_cell(money(d['G']),price_ev(d),price_basis(d))}</td><td>{ev_cell(str(d['K'])+' min',ed_status(d),ed_name(d)+'; '+ed_note(d))}</td><td class=n>{pct(d['AQ'])}</td><td class=n>{pct(d['AU'])}</td><td>{d['AX']}</td><td class=n>{d['N']}</td></tr>")
t.append('</tbody></table><p class="small">Mordialloc and Seaford are Watchlist locations, not Active Shortlist locations. Frankston is on the Active Shortlist (A3).</p>')
om=['<table><thead><tr><th>Location</th><th>Report group</th><th>Hard-filter result</th><th>Weighted score</th><th>Rank</th><th>Glen Waverley drive</th><th>Price</th><th>ED</th></tr></thead><tbody>']
for d in [x for x in rows if x['C']=='Open-mind viewing']:
    om.append(f"<tr><td>{d['A']}</td><td>{d['grp']}</td><td>{fil(d)}</td><td class=n>{wsc(d)}</td><td class=n>{rank(d)}</td><td>{ev_cell(str(d['J'])+' min',drive_ev(d),drive_basis(d))}</td><td>{ev_cell(money(d['G']),price_ev(d),price_basis(d))}</td><td>{ev_cell(str(d['K'])+' min',ed_status(d),ed_name(d)+'; '+ed_note(d))}</td></tr>")
om.append('</tbody></table><p class="small">Open-mind viewing places are not active candidates.</p>')
ref='<p>Rowville is a reference point between the two groups. It is not a candidate or a comparator and is not in the workbook. Glen Waverley drive ~15 min <span class="ev ev-e">Est.</span>; ED ~15–20 min, ED not confirmed <span class="ev ev-e">Est.</span>; price ~$1,129,555, automated valuation from search text <span class="ev ev-e">Est.</span>; India-born 3.2%, other language at home 31.6% (ABS 2021) <span class="ev ev-c">Checked</span>; band Medium. Nature: Churchill National Park partly in Rowville (Victorian Places).</p>'
sec('A5. Nature-first anchors',''.join(t)+'<h2>A6. Open-mind viewing</h2>'+''.join(om)+'<h2>A7. Reference suburb</h2>'+ref)
t=['<p class="small">Report-only groups. The workbook\'s Region column is unchanged.</p><table><thead><tr><th>Report group</th><th>Locations</th></tr></thead><tbody>']
for g in range(1,8):
    t.append(f"<tr><td>{g}. {GN[g]}</td><td>{', '.join(E(d['A'])+(' (comparator)' if is_comp(d) else '') for d in order if d['grp']==GN[g])}</td></tr>")
t.append('</tbody></table>')
p=['<p class="small">Parents\' inland/bush fit as recorded in the workbook, beside the Nature score (0–10). &ldquo;Not assessed&rdquo; is shown as recorded.</p><table><thead><tr><th>Location</th><th>Report group</th><th>Parents\' inland/bush fit</th><th>Nature score</th></tr></thead><tbody>']
for d in order: p.append(f"<tr><td>{E(d['A'])}</td><td>{E(d['grp'])}</td><td>{E(str(d['D']))}</td><td class=n>{num(d['N'])}</td></tr>")
p.append('</tbody></table>')
sec('A8. Regional groupings',''.join(t)+'<h2>A9. Parents\' inland/bushland fit</h2>'+''.join(p))
fw=[d for d in rows if str(d['L']).startswith('Fails') and d['C']!='Active Shortlist']
t=['<table><thead><tr><th>Location</th><th>Outcome</th><th>Reason</th></tr></thead><tbody>',
 '<tr><td>Point Cook</td><td><b>Conditional</b> (not a comparator; not in the workbook)</td><td>Point Cook fails the verified 90-minute Glen Waverley journey at 08:00 peak. Its inclusion would require a family decision to accept a longer travel trade-off. Peak (08:00): 55 min–1 hr 40 (top 100 min) <span class="ev ev-c">Checked</span>; 10:00 weekday 45–70 min <span class="ev ev-c">Checked</span>. Price $832,750 suburb sale median <span class="ev ev-c">Checked</span>; Werribee Mercy ED 10–16 min <span class="ev ev-c">Checked</span>.</td></tr>',
 '<tr><td>Tarneit</td><td><b>Excluded — family decision 28 Oct 2026</b></td><td>Excluded — family decision 28 Oct 2026 (stated reasons: safety concerns and not enough going for it). This is the family\'s stated view, not a checked fact. It is not a drive-time failure: the Glen Waverley drive at 10:00 on a weekday is 50–80 min <span class="ev ev-c">Checked</span>, inside the 90-min limit; peak (08:00) 1 hr–1 hr 50 is recorded for reference only.</td></tr>',
 '<tr><td>Mount Waverley</td><td><b>Excluded</b></td><td>Reported house median value $1,605,243 exceeds the $1.5M cap (OnTheHouse automated valuation, 30 Jun 2026) <span class="ev ev-e">Est.</span>.</td></tr>']
for d in fw: t.append(f"<tr><td>{E(d['A'])}</td><td><span class='st st-f'>Fail</span> (workbook status: {E(d['C'])})</td><td>Glen Waverley drive {d['J']} min ({lc(drive_basis(d))}, not peak) is over the 90-min limit.</td></tr>")
t.append('</tbody></table><p class="small">Ocean Grove and Clifton Springs also fail the drive limit but stay on the Active Shortlist (A3, group C).</p>')
oq2='''<p><b>Also carried forward</b></p><ul><li>No specific home has been checked anywhere, so no place is a Confirmed pass.</li>
<li>Glen Waverley drive times for the ranked locations were looked up with &ldquo;leave now&rdquo; on 26–27 Oct, not at peak hour. The six comparators use the top of the 10:00 weekday range (28 Oct), as in the workbook; their 08:00 peak ranges are for reference only and not used in the workbook J/K figures.</li>
<li>Living arrangement and garden are still to discuss.</li><li>Many prices are automated valuations (marked Est.), not verified sale medians.</li></ul>'''
sec('A10. Conditional and excluded locations',''.join(t)+'<h2>A11. Open questions</h2>'+OQ+oq2)
src='''<ul><li><b>Workbook:</b> &ldquo;Our Next Home — Family Relocation Scoring Model.xlsx&rdquo;, version of 29 Oct 2026 with the Glen Waverley median update (commit 5704d65; scores and ranks unchanged from 387763a). Each figure's source is recorded in its row (price basis column and notes column).</li>
<li><b>Research file:</b> &ldquo;Community comparators — research (28 Oct 2026).md&rdquo; (comparator facilities, Rowville, Point Cook, Tarneit, Mount Waverley).</li>
<li><b>Drive times:</b> Google Maps typical traffic. Ranked locations: &ldquo;leave now&rdquo;, 26–27 Oct 2026. Comparators: top of the typical range departing 10:00 on a weekday, checked 28 Oct 2026 (same basis for ED). Peak (08:00), not used in the workbook J/K figures: shown in Appendix A4 for reference. Point Cook and Tarneit: 10:00 and 08:00, checked 28 Oct 2026.</li>
<li><b>Prices:</b> as recorded in the workbook: suburb sale medians (e.g. Domain, Property.com.au, REA, OnTheHouse, Aussie/Cotality), automated valuations (OnTheHouse/Cotality; marked Est.), and Statements of Information for comparators (checked 28 Oct 2026). Wheelers Hill budget: Jellis Craig Wheelers Hill Suburb Report, Sep 2026, data to 6 Sep 2026.</li>
<li><b>Emergency departments:</b> Victorian Department of Health public-hospitals list and healthdirect (comparators, 28 Oct 2026); other rows as recorded in the workbook.</li>
<li><b>Community facilities:</b> Google Business Profiles, checked 28 Oct 2026.</li>
<li><b>Population figures:</b> ABS 2021 Census General Community Profile, Suburbs and Localities, Victoria.</li></ul>'''
sec('A12. Sources and dates',src,last=True)
H.append('</body></html>')
open(OUT,'w').write('\n'.join(H)); print('written',OUT)
