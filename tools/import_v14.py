"""Import questionnaire v14 and v15 answers (task N3) from a page `responses` export into a COPY of the workbook.

v15 (31 Oct 2026) = v14 + Timing part (R9.1-R9.4, R9 Notes). The file keeps its v14 name because docs refer to it.

Usage (from repo root):
  .venv/bin/python tools/import_v14.py <export.json> <in.xlsx> <out.xlsx>
  .venv/bin/python tools/import_v14.py <export.json> <in.xlsx> <out.xlsx> --blank-context

Never writes to <in.xlsx>. Takes, per person (Mum, Dad, Me), the latest record with pageVersion 14 or 15;
ignores TEST, other people and every other version. A v14 record leaves the R9 columns blank and prints a FLAG. New RawResponses row per person = their latest
row carried forward, overwritten by the v14 answers (rounds 5-7 design doc, section 3, Storage).
Q10 and Q13d are blanked (not asked in v14). Columns v14 never asks keep the carried-forward value
unless --blank-context is given. Exit 1 on any unmapped answer key or failed check.
"""
import argparse, datetime as dt, json, sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

PERSONS = ("Mum", "Dad", "Me")
LAST_N2_HEADER = ("BE", "[R6 Notes]")

# New RawResponses columns, appended in this order after BE.
NEW_COLUMNS = [
    ("Q3c", "[Q3c] Budget: is $1.5m right"),
    ("R8.1", "[R8.1] Kind of nature most wanted"),
    ("R8.2", "[R8.2] Kind of nature second"),
    ("R8.3-30", "[R8.3-30] Family drive ladder: 30 min step"),
    ("R8.3-20", "[R8.3-20] Family drive ladder: 20 min step"),
    ("R8.3-45", "[R8.3-45] Family drive ladder: 45 min step"),
    ("R8.3-60", "[R8.3-60] Family drive ladder: 60 min step"),
    ("R8.3-75", "[R8.3-75] Family drive ladder: 75 min step"),
    ("R8.3", "[R8.3] Family drive: extra minutes accepted (worked out)"),
    ("R8.3c", "[R8.3c] Family drive: capped by firm limit (worked out)"),
    ("R8.4", "[R8.4] Kind of area most wanted"),
    ("R8.5", "[R8.5] Kind of area least wanted"),
    ("R8.6", "[R8.6] Bushfire risk tolerance (feeds nothing)"),
    # v15 Timing part: feeds no weight, filter or score (market-timing note only).
    ("R9.1", "[R9.1] Timing: when to have moved by"),
    ("R9.2", "[R9.2] Timing: if prices fell 5% after buying"),
    ("R9.3", "[R9.3] Timing: sell Beachmere first or buy first"),
    ("R9.4", "[R9.4] Timing: where to live until we buy"),
    ("R9 Notes", "[R9 Notes] Timing: notes"),
]

# Every tag the v14 page can store, except notes that have no column of their own.
V14_TAGS = ["Q7a", "Q3c", "R5.2", "R5.4", "R8.1", "R8.2", "R8.4", "R8.5", "R8.6",
            "R8.3-30", "R8.3-20", "R8.3-45", "R8.3-60", "R8.3-75", "R8.3", "R8.3c",
            "Q11", "Q12", "Q13", "R5.5", "R5.6", "R5.7", "R5 Notes",
            "R6.1", "R6.2", "R6.3", "R6.4", "R6.5a", "R6.5b", "R6.6", "R6 Notes",
            "Q14"] + [f"Q15.{i}" for i in range(1, 11)] + ["Q15 Notes"]
R9_TAGS = ["R9.1", "R9.2", "R9.3", "R9.4", "R9 Notes"]
VERSIONS = (14, 15)
Q15_TAGS = {f"Q15.{i}" for i in range(1, 11)}
NOTES_TO_AM = "Q10-Q13c Notes"  # no RawResponses column; kept as a labelled note in the AM merge log
BLANKED = ["Q10", "Q13d"]       # round 8 note: not asked in v14, columns stay empty for v14
# Columns v14 never asks: carried forward by default, blanked with --blank-context.
CONTEXT_TAGS = ["Q1", "Q2", "Q3", "Q4", "Q4b", "Q5", "Q6", "Q7", "Q7b", "Q8", "Q9", "Q13b", "Q13c",
                "R5.1", "R5.3", "R6.7", "Q16", "Q17", "Q18", "Q19", "Q20", "Q21"]
R86_X = "X — Not asked (most-wanted area wasn't B or C)"

CA_LABEL_ROW = 73  # CleanedAnswers: appended below the derived block (rows 58-71), never inserted


def tag_of(header):
    return header[1:header.index("]")] if isinstance(header, str) and header.startswith("[") and "]" in header else None


def select(records):
    v14 = [r for r in records if r.get("pageVersion") in VERSIONS and r.get("person") in PERSONS]
    chosen, others = {}, []
    for r in sorted(v14, key=lambda r: r["savedAt"]):
        if r["person"] in chosen:
            others.append(chosen[r["person"]])
        chosen[r["person"]] = r
    skipped = [r for r in records if r not in v14]
    return chosen, others, skipped


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("export")
    ap.add_argument("inp")
    ap.add_argument("out")
    ap.add_argument("--blank-context", action="store_true", help="blank the columns v14 never asks instead of carrying them forward")
    a = ap.parse_args()
    if Path(a.inp).resolve() == Path(a.out).resolve():
        sys.exit("out must differ from in: this tool never writes the input workbook")

    records = json.loads(Path(a.export).read_text(encoding="utf-8"))["records"]
    chosen, others, skipped = select(records)
    print(f"records: {len(records)}; v14/v15 Mum/Dad/Me: {len(chosen) + len(others)}; ignored (other version/person/TEST): {len(skipped)}")
    for r in others:
        print(f"  superseded v{r['pageVersion']} record (not imported): {r['id']} {r['person']} {r['savedAt']}")
    if not chosen:
        sys.exit("no v14/v15 records to import")

    known = set(V14_TAGS) | set(R9_TAGS) | {NOTES_TO_AM}
    for r in chosen.values():
        bad = sorted(set(r["answers"]) - known)
        if bad:
            sys.exit(f"{r['id']}: unmapped answer keys {bad}")
        if r.get("answeredCount") != len(r["answers"]):
            sys.exit(f"{r['id']}: answeredCount {r.get('answeredCount')} != {len(r['answers'])} answers")

    wb = openpyxl.load_workbook(a.inp)
    ws = wb["RawResponses"]
    if ws[f"{LAST_N2_HEADER[0]}1"].value != LAST_N2_HEADER[1]:
        sys.exit(f"RawResponses {LAST_N2_HEADER[0]}1 is {ws[LAST_N2_HEADER[0] + '1'].value!r}, expected {LAST_N2_HEADER[1]!r}")

    cols = {tag_of(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1) if tag_of(ws.cell(1, c).value)}
    src = ws[f"{LAST_N2_HEADER[0]}1"]
    for tag, header in NEW_COLUMNS:
        if tag in cols:
            continue
        c = ws.max_column + 1
        cell = ws.cell(1, c, header)
        cell._style = copy(src._style)
        ws.column_dimensions[get_column_letter(c)].width = ws.column_dimensions[LAST_N2_HEADER[0]].width
        cols[tag] = c
        print(f"  added RawResponses {get_column_letter(c)}1 {header}")
    missing = [t for t in V14_TAGS + BLANKED + CONTEXT_TAGS + ["Gap-form notes"] if t not in cols]
    if missing:
        sys.exit(f"RawResponses has no column for {missing}")
    last_col = get_column_letter(ws.max_column)

    flags = []
    for person in PERSONS:
        r = chosen.get(person)
        if not r:
            continue
        ans, ts = r["answers"], r["savedAt"][:19]
        prev = max(row for row in range(2, ws.max_row + 1) if ws.cell(row, 2).value == person)
        am_prev = ws.cell(prev, cols["Gap-form notes"]).value or ""
        ver = r["pageVersion"]
        if f"v{ver} {ts}Z" in am_prev:
            sys.exit(f"{person}: record {ts}Z already imported (row {prev})")
        new = ws.max_row + 1
        for c in range(1, ws.max_column + 1):
            s, d = ws.cell(prev, c), ws.cell(new, c)
            d.value = s.value
            d._style = copy(s._style)

        ws.cell(new, 1).value = dt.datetime.fromisoformat(ts)
        for tag in V14_TAGS + R9_TAGS:
            v = ans.get(tag)
            if v is not None and tag in Q15_TAGS:
                v = v[0]
            ws.cell(new, cols[tag]).value = v
        for tag in BLANKED:
            ws.cell(new, cols[tag]).value = None
        if a.blank_context:
            for tag in CONTEXT_TAGS:
                ws.cell(new, cols[tag]).value = None

        log = f"v{ver} {ts}Z merged onto row {prev} (record {r['id']}; Q10, Q13d blanked; context columns {'blanked' if a.blank_context else 'carried forward'})."
        if "R8.6" not in ans:
            area = ans.get("R8.4", "")[:1]
            if area in ("A", "D", "E"):
                ws.cell(new, cols["R8.6"]).value = R86_X
                log += " R8.6 backfilled X (answered before R8.6 existed)."
            else:
                flags.append(f"{person} row {new}: R8.6 Needs data: answered before R8.6 existed (R8.4 = {area or 'blank'})")
                log += " R8.6 Needs data: answered before R8.6 existed."
        if ver == 14:
            flags.append(f"{person} row {new}: R9.1-R9.4 Needs data: answered v14 (no Timing part); add chat answers by hand, source 'chat'")
            log += " R9 Needs data: v14 record (no Timing part)."
        if ans.get(NOTES_TO_AM):
            log += f" Q10-Q13c notes: {ans[NOTES_TO_AM]}"
        ws.cell(new, cols["Gap-form notes"]).value = f"{log} | {am_prev}" if am_prev else log
        print(f"  {person}: row {new} = row {prev} + v{ver} record {r['id']} ({len(ans)} answers)")

    ca = wb["CleanedAnswers"]
    have = {ca.cell(row, 1).value for row in range(1, ca.max_row + 1)}
    pattern = ca["B41"].value
    if "$A:$BE" not in pattern or "$A$1:$BE$1" not in pattern:
        sys.exit("CleanedAnswers B41 no longer follows the N2 lookup pattern")
    todo = [h for _, h in NEW_COLUMNS if h not in have]
    if todo:
        if ca.max_row >= CA_LABEL_ROW:
            sys.exit(f"CleanedAnswers has content at or below row {CA_LABEL_ROW}; place new rows by hand")
        ca.cell(CA_LABEL_ROW, 1, "v14/v15 answers (round 8 + timing; context, not weighted)")._style = copy(ca["A30"]._style)
        for i, h in enumerate(["Question ID", "Mum Answer", "Dad Answer", "Me Answer", "Mum Code", "Dad Code", "Me Code"]):
            ca.cell(CA_LABEL_ROW + 1, i + 1, h)._style = copy(ca.cell(31, i + 1)._style)
        for i, h in enumerate(todo):
            row = CA_LABEL_ROW + 2 + i
            ca.cell(row, 1, h)._style = copy(ca["A41"]._style)
            for j, col in enumerate("BCD"):
                f = ca[f"{col}41"].value.replace("$A:$BE", f"$A:${last_col}").replace("$A$1:$BE$1", f"$A$1:${last_col}$1").replace("$A41", f"$A{row}")
                ca[f"{col}{row}"] = f
                ca[f"{col}{row}"]._style = copy(ca[f"{col}41"]._style)
                ca[f"{'EFG'[j]}{row}"] = f"=LEFT({col}{row},1)"
                ca[f"{'EFG'[j]}{row}"]._style = copy(ca[f"{'EFG'[j]}41"]._style)
        print(f"  added CleanedAnswers rows {CA_LABEL_ROW}-{CA_LABEL_ROW + 1 + len(todo)} (lookups over RawResponses $A:${last_col})")

    for f in flags:
        print("FLAG", f)
    wb.save(a.out)
    print(f"saved {a.out}")


if __name__ == "__main__":
    main()
